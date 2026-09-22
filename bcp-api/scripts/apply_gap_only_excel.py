"""Update ONLY gap_description in cbuae-aml-demo-judgments-94.json from the latest
gap-analysis Excel, for clauses whose Compliance Status is Partial or Non-Compliant.

No other field (status, interpretation, confidence, policy_extract, document_reference,
suggested_action, gap_direction, clause content) is touched.

After running this, run sync_demo_seed_94.py to propagate to the live seed file
(cbuae-aml-demo-judgments.json) and the frontend clause list.
"""
import json
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SEED_94 = ROOT / "SeedData" / "cbuae-aml-demo-judgments-94.json"
EXCEL = Path(r"C:\Users\Pc\Downloads\Tester-4_gap_analysis 2090226 (1).xlsx")


def normalize_key(value: str) -> str:
    trimmed = (value or "").strip()
    while trimmed.endswith("."):
        trimmed = trimmed[:-1]
    return trimmed


def load_excel_gap_map(excel_path: Path) -> dict[str, str]:
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb["Gap Analysis"]
    mapping: dict[str, str] = {}
    for row in ws.iter_rows(min_row=11, max_row=ws.max_row, values_only=True):
        clause = row[1]
        if clause is None or str(clause).strip() == "":
            continue
        status = (row[13] or "").strip()
        if status not in ("Partial", "Non-Compliant"):
            continue
        gap_text = (row[3] or "").strip() if row[3] else ""
        if not gap_text:
            continue
        mapping[normalize_key(str(clause))] = gap_text
    wb.close()
    return mapping


def main():
    gap_map = load_excel_gap_map(EXCEL)
    print(f"Excel gap rows (Partial/Non-Compliant with text): {len(gap_map)}")

    seed = json.loads(SEED_94.read_text(encoding="utf-8"))
    updated = 0
    no_excel_match = []
    for item in seed:
        key = normalize_key(item["clause_no"])
        if key not in gap_map:
            if item.get("overall_status") in ("partial", "non-compliant"):
                no_excel_match.append(item["clause_no"])
            continue
        new_gap = gap_map[key]
        if item.get("gap_description") != new_gap:
            item["gap_description"] = new_gap
            updated += 1

    SEED_94.write_text(json.dumps(seed, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")

    print(f"Updated gap_description on {updated} clauses in {SEED_94.name}")
    if no_excel_match:
        print(f"WARNING: {len(no_excel_match)} partial/non-compliant clauses had no Excel gap match: {no_excel_match}")


if __name__ == "__main__":
    main()
